"""Safe, bounded text extraction for user-provided chat documents."""

from __future__ import annotations

import csv
import datetime as dt
import io
import math
import re
import zipfile
from contextlib import contextmanager
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from agents.agent_harness.redaction import is_sensitive_key


MIB = 1024 * 1024
MAX_DOCUMENT_BYTES = 20 * MIB
MAX_DOCUMENT_TEXT_CHARS = 24_000
MAX_PDF_PAGES = 60
MAX_XLSX_SHEETS = 20
MAX_SHEET_ROWS = 1_000
MAX_SHEET_COLUMNS = 80
MAX_IMPORT_PREVIEW_ROWS = 100
MAX_ANALYSIS_ROWS = 10_000
MAX_XLSX_ARCHIVE_ENTRIES = 5_000
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * MIB
MAX_XLSX_COMPRESSION_RATIO = 500

CHAT_DOCUMENT_TYPES = {
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ),
}
_TEXT_TYPES = frozenset({"text/csv", "text/markdown", "text/plain"})
_XLSX_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_MIME_COMPATIBILITY = {
    "text/csv": frozenset({
        "text/csv",
        "application/csv",
        "text/plain",
        "application/octet-stream",
    }),
    "text/markdown": frozenset({
        "text/markdown",
        "text/x-markdown",
        "text/plain",
        "application/octet-stream",
    }),
    "text/plain": frozenset({
        "text/plain",
        "application/octet-stream",
    }),
    "application/pdf": frozenset({
        "application/pdf",
        "application/octet-stream",
    }),
    _XLSX_TYPE: frozenset({
        _XLSX_TYPE,
        "application/vnd.ms-excel",
        "application/octet-stream",
    }),
}
_MIME_BY_EXTENSION = {
    ".gif": "image/gif",
    ".jpeg": "image/jpeg",
    ".jpg": "image/jpeg",
    ".mp4": "video/mp4",
    ".mov": "video/quicktime",
    ".png": "image/png",
    ".webm": "video/webm",
    ".webp": "image/webp",
}
IMAGE_EXTENSIONS = frozenset({".gif", ".jpeg", ".jpg", ".png", ".webp"})
VIDEO_EXTENSIONS = frozenset({".mp4", ".mov", ".webm"})


def detect_chat_attachment_type(
    content: bytes,
    *,
    filename: str,
    content_type: str,
    media_signatures: dict[str, Any],
) -> str:
    """Resolve a supported MIME type from a safe extension and file signature."""
    extension = Path(filename).suffix.lower()
    requested_type = str(content_type or "").split(";", 1)[0].strip().lower()
    expected_media_type = _MIME_BY_EXTENSION.get(extension)
    expected_document_type = CHAT_DOCUMENT_TYPES.get(extension)
    if extension == ".xls":
        raise ValueError("暂不支持旧版 .xls，请另存为 .xlsx 后上传")
    if expected_media_type is None and expected_document_type is None:
        expected_media_type = (
            requested_type if requested_type in media_signatures else None
        )
    if expected_media_type:
        if requested_type and requested_type not in {
            expected_media_type,
            "application/octet-stream",
        }:
            raise ValueError("attachment filename and media type do not match")
        if not media_signatures[expected_media_type](content):
            raise ValueError("attachment content does not match its file type")
        return expected_media_type

    if expected_document_type is None:
        raise ValueError(
            "支持的附件格式：图片、视频、CSV、XLSX、PDF、Markdown 和 TXT"
        )
    if requested_type and requested_type not in _MIME_COMPATIBILITY[
        expected_document_type
    ]:
        raise ValueError("attachment filename and media type do not match")
    if expected_document_type == "application/pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("attachment content does not match its file type")
    elif expected_document_type == _XLSX_TYPE:
        _validate_xlsx_archive(content)
    else:
        _decode_text(content)
    return expected_document_type


def extract_chat_document_text(
    content: bytes,
    *,
    filename: str,
    media_type: str,
) -> str:
    """Extract a bounded text representation without executing file content."""
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("文档附件超过 20 MiB 限制")
    if media_type in _TEXT_TYPES:
        text = _decode_text(content)
        if media_type == "text/csv":
            return _extract_csv(text)
        return _truncate(text.strip(), MAX_DOCUMENT_TEXT_CHARS)
    if media_type == _XLSX_TYPE:
        return _extract_xlsx(content)
    if media_type == "application/pdf":
        return _extract_pdf(content)
    raise ValueError(f"不支持解析文档类型：{media_type}")


def read_chat_document_table(
    content: bytes,
    *,
    filename: str,
    media_type: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Read a bounded CSV/XLSX table for a non-executing import preview."""
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("文档附件超过 20 MiB 限制")
    extension = Path(filename).suffix.lower()
    if extension == ".csv" and media_type == "text/csv":
        return _read_csv_table(content)
    if extension == ".xlsx" and media_type == _XLSX_TYPE:
        return _read_xlsx_table(content, sheet_name=sheet_name)
    raise ValueError("广告配置导入仅支持 CSV 和 XLSX 文件")


@contextmanager
def open_chat_document_table(
    content: bytes,
    *,
    filename: str,
    media_type: str,
    sheet_name: str | None = None,
) -> Iterator[tuple[list[str], Iterator[list[str]], str]]:
    """Stream one bounded CSV/XLSX sheet without exposing the whole file to the LLM."""
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("文档附件超过 20 MiB 限制")
    extension = Path(filename).suffix.lower()
    if extension == ".csv" and media_type == "text/csv":
        if sheet_name:
            raise ValueError("CSV 文件没有工作表")
        reader = csv.reader(io.StringIO(_decode_text(content), newline=""))
        try:
            headers = _validate_table_headers(next(reader, ()) or ())
        except csv.Error as error:
            raise ValueError("CSV 格式无效，无法读取表格") from error
        yield headers, _bounded_csv_rows(reader, len(headers)), ""
        return
    if extension != ".xlsx" or media_type != _XLSX_TYPE:
        raise ValueError("表格分析仅支持 CSV 和 XLSX 文件")
    with _open_xlsx_analysis(content, sheet_name) as table:
        yield table


def _bounded_csv_rows(reader: Any, column_count: int) -> Iterator[list[str]]:
    try:
        for index, row in enumerate(reader):
            if index >= MAX_ANALYSIS_ROWS:
                raise ValueError("表格超过 10000 行分析上限，无法计算全量统计")
            if len(row) > column_count:
                raise ValueError("CSV 数据行列数超过表头")
            yield [str(cell).strip() for cell in row] + [""] * (column_count - len(row))
    except csv.Error as error:
        raise ValueError("CSV 格式无效，无法读取表格") from error


def _bounded_xlsx_rows(iterator: Any, worksheet: Any) -> Iterator[list[str]]:
    for index, row in enumerate(iterator):
        if index >= MAX_ANALYSIS_ROWS:
            raise ValueError("表格超过 10000 行分析上限，无法计算全量统计")
        yield [_table_cell(value) for value in row]
    if (worksheet.max_row or 0) > MAX_ANALYSIS_ROWS + 1:
        raise ValueError("表格超过 10000 行分析上限，无法计算全量统计")


@contextmanager
def _open_xlsx_analysis(
    content: bytes, sheet_name: str | None,
) -> Iterator[tuple[list[str], Iterator[list[str]], str]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    _validate_xlsx_archive(content)
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False,
        )
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise ValueError("XLSX 最多支持 20 个工作表")
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError("指定的 XLSX 工作表不存在")
            worksheet = workbook[sheet_name]
        elif workbook.worksheets:
            worksheet = workbook.worksheets[0]
        else:
            raise ValueError("XLSX 中没有可读取的工作表")
        if (worksheet.max_column or 0) > MAX_SHEET_COLUMNS:
            raise ValueError("XLSX 列数超过 80 列限制")
        iterator = worksheet.iter_rows(
            max_row=MAX_ANALYSIS_ROWS + 2,
            max_col=min(worksheet.max_column or MAX_SHEET_COLUMNS, MAX_SHEET_COLUMNS),
            values_only=True,
        )
        headers = _validate_table_headers(next(iterator, ()) or ())
        yield headers, _bounded_xlsx_rows(iterator, worksheet), str(worksheet.title)
    except (InvalidFileException, KeyError, OSError, zipfile.BadZipFile) as error:
        raise ValueError("无法读取 XLSX 工作簿，请检查文件是否损坏") from error
    finally:
        if workbook is not None:
            workbook.close()


def _read_csv_table(content: bytes) -> dict[str, Any]:
    text = _decode_text(content)
    try:
        reader = csv.reader(io.StringIO(text, newline=""))
        header = next(reader, None)
        headers = _validate_table_headers(header or [])
        rows: list[list[str]] = []
        truncated = False
        for row in reader:
            if not any(str(cell).strip() for cell in row):
                continue
            if len(row) > MAX_SHEET_COLUMNS:
                raise ValueError("CSV 列数超过 80 列限制")
            if len(rows) >= MAX_IMPORT_PREVIEW_ROWS:
                truncated = True
                break
            rows.append([str(cell).strip() for cell in row])
    except csv.Error as error:
        raise ValueError("CSV 格式无效，无法读取表格") from error
    return {
        "headers": headers,
        "rows": rows,
        "sheet_name": "",
        "truncated": truncated,
    }


def _read_xlsx_table(
    content: bytes,
    *,
    sheet_name: str | None,
) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    _validate_xlsx_archive(content)
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise ValueError("XLSX 最多支持 20 个工作表")
        if sheet_name:
            worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else None
            if worksheet is None:
                raise ValueError("指定的 XLSX 工作表不存在")
        elif workbook.worksheets:
            worksheet = workbook.worksheets[0]
        else:
            raise ValueError("XLSX 中没有可读取的工作表")

        if (worksheet.max_column or 0) > MAX_SHEET_COLUMNS:
            raise ValueError("XLSX 列数超过 80 列限制")
        max_row = min(worksheet.max_row or 0, MAX_IMPORT_PREVIEW_ROWS + 1)
        iterator = worksheet.iter_rows(
            max_row=max_row,
            max_col=min(worksheet.max_column or MAX_SHEET_COLUMNS, MAX_SHEET_COLUMNS),
            values_only=True,
        )
        headers = _validate_table_headers(next(iterator, ()) or ())
        rows: list[list[str]] = []
        for row in iterator:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(rows) >= MAX_IMPORT_PREVIEW_ROWS:
                break
            rows.append([_table_cell(value) for value in row])
        truncated = (worksheet.max_row or 0) > MAX_IMPORT_PREVIEW_ROWS + 1
        return {
            "headers": headers,
            "rows": rows,
            "sheet_name": str(worksheet.title),
            "truncated": truncated,
        }
    except ValueError:
        raise
    except (InvalidFileException, KeyError, OSError, zipfile.BadZipFile) as error:
        raise ValueError("无法读取 XLSX 工作簿，请检查文件是否损坏") from error
    finally:
        if workbook is not None:
            workbook.close()


def _validate_table_headers(values: Any) -> list[str]:
    if not isinstance(values, (list, tuple)) or not values:
        raise ValueError("表格 header 不能为空")
    if len(values) > MAX_SHEET_COLUMNS:
        raise ValueError("表格列数超过 80 列限制")
    headers = [str(value or "").strip() for value in values]
    normalized = [header.casefold() for header in headers]
    if any(not header for header in headers) or len(set(normalized)) != len(headers):
        raise ValueError("表格 header 必须非空且不能重复")
    if any(len(header) > 200 for header in headers):
        raise ValueError("表格 header 单元格不能超过 200 个字符")
    _reject_sensitive_headers(headers)
    return headers


def _reject_sensitive_headers(headers: Any) -> None:
    if any(is_sensitive_key(header) for header in headers):
        raise ValueError("表格包含敏感凭证列，不能用于聊天分析或预览")


def _table_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("XLSX 表格包含非有限数值")
        return str(value)
    return str(value).strip()


def _decode_text(content: bytes) -> str:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("文本文件需使用 UTF-8 编码") from error
    if "\x00" in text:
        raise ValueError("文本文件包含不支持的二进制内容")
    return "".join(
        character
        for character in text
        if character in "\t\n\r" or ord(character) >= 32
    )


def _extract_csv(text: str) -> str:
    output: list[str] = []
    try:
        rows = csv.reader(io.StringIO(text, newline=""))
        for row_index, row in enumerate(rows):
            if row_index >= MAX_SHEET_ROWS:
                output.append("[内容已截断：最多读取 1000 行]")
                break
            if row_index == 0:
                _reject_sensitive_headers(row)
            cells = [
                " ".join(str(cell).replace("|", "\\|").split())[:500]
                for cell in row[:MAX_SHEET_COLUMNS]
            ]
            if row_index == 0 and len(row) > MAX_SHEET_COLUMNS:
                cells.append("[其余列已省略]")
            output.append(" | ".join(cells))
    except csv.Error as error:
        raise ValueError("CSV 格式无效，无法读取表格") from error
    return _truncate("\n".join(output), MAX_DOCUMENT_TEXT_CHARS)


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except (OSError, zipfile.BadZipFile) as error:
        raise ValueError("XLSX 文件结构无效") from error
    if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
        raise ValueError("XLSX 文件包含过多内部条目")
    total_size = 0
    for entry in entries:
        if entry.is_dir():
            continue
        if entry.flag_bits & 0x1:
            raise ValueError("加密 XLSX 文件暂不支持")
        total_size += entry.file_size
        if (
            entry.file_size > 0
            and entry.file_size / max(entry.compress_size, 1)
            > MAX_XLSX_COMPRESSION_RATIO
        ):
            raise ValueError("XLSX 压缩比例异常，已拒绝解析")
    if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX 解压后超过 64 MiB 安全上限")
    required = {"[Content_Types].xml", "xl/workbook.xml"}
    if not required.issubset({entry.filename for entry in entries}):
        raise ValueError("XLSX 文件缺少工作簿结构")


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    _validate_xlsx_archive(content)
    workbook = None
    output: list[str] = []
    truncated = False
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise ValueError("XLSX 最多支持 20 个工作表")
        remaining = MAX_DOCUMENT_TEXT_CHARS
        for worksheet in workbook.worksheets:
            if remaining <= 0:
                truncated = True
                break
            output.append(f"工作表：{worksheet.title}")
            max_row = min(worksheet.max_row or 0, MAX_SHEET_ROWS)
            max_column = min(worksheet.max_column or 0, MAX_SHEET_COLUMNS)
            if (worksheet.max_row or 0) > MAX_SHEET_ROWS:
                truncated = True
            if (worksheet.max_column or 0) > MAX_SHEET_COLUMNS:
                truncated = True
            rows = worksheet.iter_rows(
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            )
            for row_index, row in enumerate(rows):
                if row_index == 0:
                    _reject_sensitive_headers(row)
                cells = [
                    _xlsx_cell(value)
                    for value in row
                ]
                line = " | ".join(cells).rstrip()
                if not line:
                    continue
                output.append(line)
                remaining -= len(line) + 1
                if remaining <= 0:
                    truncated = True
                    break
        if truncated:
            output.append("[内容已截断：受行数、列数或文本长度限制]")
    except (InvalidFileException, KeyError, OSError, zipfile.BadZipFile) as error:
        raise ValueError("无法读取 XLSX 工作簿，请检查文件是否损坏") from error
    finally:
        if workbook is not None:
            workbook.close()
    extracted = "\n".join(output).strip()
    if not extracted:
        raise ValueError("XLSX 中没有可读取的单元格内容")
    return _truncate(extracted, MAX_DOCUMENT_TEXT_CHARS)


def _xlsx_cell(value: Any) -> str:
    if value is None:
        return ""
    cleaned = re.sub(r"[\x00-\x1f\x7f]", " ", str(value))
    return " ".join(cleaned.replace("|", "\\|").split())[:500]


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ValueError("加密 PDF 暂不支持，请先解除密码保护")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF 最多支持 60 页")
        output: list[str] = []
        remaining = MAX_DOCUMENT_TEXT_CHARS
        for index, page in enumerate(reader.pages):
            if remaining <= 0:
                output.append("[内容已截断：达到文本长度上限]")
                break
            text = page.extract_text() or ""
            normalized = "\n".join(
                " ".join(line.split())
                for line in text.splitlines()
                if line.strip()
            )
            if normalized:
                section = f"[第 {index + 1} 页]\n{normalized}"
                output.append(section[:remaining])
                remaining -= len(section) + 1
    except ValueError:
        raise
    except (PdfReadError, EOFError, KeyError, OSError, TypeError) as error:
        raise ValueError("无法读取 PDF，请检查文件是否损坏") from error
    extracted = "\n".join(output).strip()
    if not extracted:
        raise ValueError("PDF 中没有可提取文本，扫描版 PDF 暂不支持 OCR")
    return _truncate(extracted, MAX_DOCUMENT_TEXT_CHARS)


def _truncate(value: str, limit: int) -> str:
    normalized = value[:limit].strip()
    if len(value) > limit:
        suffix = "\n[内容已截断]"
        normalized = normalized[: max(0, limit - len(suffix))] + suffix
    return normalized


__all__ = [
    "CHAT_DOCUMENT_TYPES",
    "IMAGE_EXTENSIONS",
    "MAX_IMPORT_PREVIEW_ROWS",
    "MAX_DOCUMENT_BYTES",
    "MAX_DOCUMENT_TEXT_CHARS",
    "MAX_ANALYSIS_ROWS",
    "VIDEO_EXTENSIONS",
    "detect_chat_attachment_type",
    "extract_chat_document_text",
    "open_chat_document_table",
    "read_chat_document_table",
]
